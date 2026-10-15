from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from core.insights_engine import generate_insights
from core.report_generator import generate_html_report, generate_pdf_report, generate_excel_report

def sample_df():
    return pd.DataFrame({"x":[1,2,3,4,5,100,7,8,9,10], "y":[2,4,6,8,10,12,14,16,18,20], "group":["A"]*8+["B"]*2})

def test_insights_have_observed_evidence():
    insights = generate_insights(sample_df())
    assert insights
    assert all("evidence" in item and "finding" in item for item in insights)

def test_all_report_formats_generate_with_real_dataset():
    df = sample_df()
    html = generate_html_report(df)
    pdf = generate_pdf_report(df)
    xlsx = generate_excel_report(df)
    assert b"Data Science Analytics Report" in html
    assert pdf[:4] == b"%PDF"
    workbook = load_workbook(BytesIO(xlsx), read_only=True)
    assert "Overview" in workbook.sheetnames
    assert "Analyzed_Data" in workbook.sheetnames
    assert "Insights" in workbook.sheetnames
    assert "Charts" in workbook.sheetnames
    workbook.close()

def test_excel_export_escapes_formula_like_untrusted_text():
    df = pd.DataFrame({"text":["=1+1", "normal"], "value":[1, -2]})
    content = generate_excel_report(df, {"profile": __import__("core.profiler", fromlist=["profile_dataset"]).profile_dataset(df),
                                        "descriptive": __import__("core.profiler", fromlist=["descriptive_table"]).descriptive_table(df),
                                        "insights": [], "chart_columns": []})
    workbook = load_workbook(BytesIO(content), read_only=False, data_only=False)
    cell = workbook["Analyzed_Data"]["A2"]
    assert cell.value == "'=1+1"
    assert cell.data_type != "f"
    workbook.close()

def test_empty_dataframe_can_still_be_reported():
    df = pd.DataFrame()
    assert generate_html_report(df)
    assert generate_pdf_report(df)[:4] == b"%PDF"
    content = generate_excel_report(df)
    workbook = load_workbook(BytesIO(content), read_only=True)
    assert "Overview" in workbook.sheetnames
    workbook.close()


def test_csv_export_escapes_formula_like_text():
    from utils.exports import safe_csv_bytes
    df = pd.DataFrame({"comment":["=1+1", "safe"], "amount":[-2, 3]})
    exported = safe_csv_bytes(df).decode("utf-8-sig")
    assert "'=1+1" in exported
    assert "-2" in exported
